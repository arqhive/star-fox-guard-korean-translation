"""번역 작업 파이프라인.
  python tl_tool.py split          -> translation/work/in_NN.json (일본어가 있는 문단만)
  python tl_tool.py check NN|all   -> translation/work/out_NN.json 검사 (태그/줄수/줄폭/문자)
  python tl_tool.py merge          -> out_*.json 을 sfg_text.xlsx '한국어 번역' 열에 반영

줄폭: 원본 MCD 글리프 폭 + 줄 간격값(a) + 커닝으로 원문 각 줄 픽셀 폭 계산.
한글 추정폭 = 폰트별 한자·가나 중앙값 - 3, 영숫자·기호 = 같은 폰트 원본 글리프 폭(없으면 한글폭 x 0.55),
공백 = 폰트표 폭 + a, 버튼 아이콘 = 원문과 같게.
한도 = max(원문 최장줄 x 1.08, 원문 최장줄 + 한글 1자) (문단 단위)
"""
import sys, os, re, json, glob, statistics
from collections import defaultdict
import openpyxl

from text_tool import XLSX, iter_mcds, para_text
from mcd_lib import parse_mcd

WORK = os.path.join(os.path.dirname(XLSX), 'work')
JP_RE = re.compile(r'[぀-ヿ㐀-鿿！-～]')
TAG_RE = re.compile(r'\{(btn|k|sp|f):(-?\d+)\}')
ALLOWED_RE = re.compile(r'^[가-힣ㄱ-ㆎ\x20-\x7a\x7c\x7e\n・…「」『』～♪★☆○×→←↑↓※·“”‘’△▽□◎●■▲▼◆◇]*$')

CHUNKS = [
    ('01', lambda r: r['mcd'] == 'messcore.mcd' and r['msg'] < 210),
    ('02', lambda r: r['mcd'] == 'messcore.mcd' and r['msg'] >= 210),
    ('03', lambda r: r['mcd'] in ('messmsg_firsttime.mcd', 'messmsg_1stage.mcd', 'messmsg_2stage.mcd', 'messmsg_3stage.mcd',
                                  'messmsg_4stage.mcd', 'messmsg_5stage.mcd', 'messmsg_epilogue.mcd', 'messevent.mcd',
                                  'messcredit.mcd')),
    ('04', lambda r: r['mcd'] in ('messmsg_jamming.mcd', 'messmsg_permissions.mcd', 'messtips.mcd', 'messprofile.mcd')),
    ('05', lambda r: r['mcd'] in ('messhud.mcd', 'messedit.mcd')),
    ('06', lambda r: True),  # title/select/result/option/loading/common
]


class Metrics:
    def __init__(self):
        self.cjk = defaultdict(list)
        self.char_w = {}  # (font, char) -> w
        self.font_tab = {}
        self.mcds = {}
        for p, dat, mcd, data in iter_mcds():
            M = parse_mcd(data)
            self.mcds[(dat, mcd)] = M
            for f in M['fonts']:
                self.font_tab[f[0]] = f
            for fo, ch, gi in M['syms']:
                w = M['glyphs'][gi][5]
                if 0x4E00 <= ch <= 0x9FFF or 0x3041 <= ch <= 0x30FF:
                    self.cjk[fo].append(w)
                self.char_w.setdefault((fo, ch), w)

    def hangul_w(self, font):
        if self.cjk.get(font):
            return statistics.median(self.cjk[font]) - 3
        return self.font_tab[font][2] * 0.75

    def jp_line_widths(self, M, pa):
        out = []
        for l in pa['lines']:
            w = l['words']; x = 0; i = 0
            while i < len(w) and w[i] != 0x8000:
                c, arg = w[i], w[i + 1]
                if c < 0x8000:
                    x += M['glyphs'][M['syms'][c][2]][5] + l['a'] + (arg - 0x10000 if arg & 0x8000 else arg)
                elif c == 0x8001:
                    x += self.font_tab[arg][1] + l['a']
                elif c == 0x8003:
                    x += self.font_tab[pa['font']][2] * 0.8
                i += 2
            out.append(x)
        return out

    def ko_line_width(self, text, font, a):
        x = 0
        hw = self.hangul_w(font)
        for m in re.finditer(r'\{btn:\d+\}|\{\{|\}\}|.', text):
            t = m.group()
            if t.startswith('{btn'):
                x += self.font_tab[font][2] * 0.8
            elif t == ' ':
                x += self.font_tab[font][1] + a
            elif '가' <= t <= '힣':
                x += hw + a
            else:
                ch = t[0]
                x += self.char_w.get((font, ord(ch)), hw * 0.55) + a
        return x


def rows_from_xlsx():
    ws = openpyxl.load_workbook(XLSX).active
    for r in ws.iter_rows(min_row=2, values_only=True):
        yield dict(id=r[0], dat=r[1], mcd=r[2], msg=r[3], para=r[4], font=r[5], event=r[6], ja=r[7], ko=r[8])


def split():
    os.makedirs(WORK, exist_ok=True)
    met = Metrics()
    chunks = defaultdict(list)
    for r in rows_from_xlsx():
        if not JP_RE.search(r['ja'].replace('{k:', '')):
            continue
        M = met.mcds[(r['dat'], r['mcd'])]
        pa = M['msgs'][r['msg']]['paras'][r['para']]
        widths = met.jp_line_widths(M, pa)
        mx = max(widths)
        limit = max(mx * 1.08, mx + met.hangul_w(r['font']))
        item = dict(id=r['id'], file=r['mcd'], event=r['event'],
                    ja=TAG_RE.sub(lambda m: m.group() if m.group(1) == 'btn' else '', r['ja']),
                    lines=len(pa['lines']), max_chars=int(limit // (met.hangul_w(r['font']) + pa['lines'][0]['a'])))
        for name, pred in CHUNKS:
            if pred(r):
                chunks[name].append(item); break
    for name, items in sorted(chunks.items()):
        json.dump(items, open(os.path.join(WORK, f'in_{name}.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        print(name, len(items), sum(len(i['ja']) for i in items), 'chars')


def check(name, met=None, quiet=False):
    met = met or Metrics()
    src = {i['id']: i for i in json.load(open(os.path.join(WORK, f'in_{name}.json'), encoding='utf-8'))}
    outs = {}
    for p in sorted(glob.glob(os.path.join(WORK, f'out_{name}*.json'))):
        for it in json.load(open(p, encoding='utf-8')):
            outs[it['id']] = it['ko']
    rows = {r['id']: r for r in rows_from_xlsx()}
    errs = []
    missing = [i for i in src if i not in outs]
    if missing:
        errs.append(f'누락 {len(missing)}개: {missing[:10]}')
    for i, ko in outs.items():
        if i not in src:
            errs.append(f'#{i}: in 파일에 없는 id'); continue
        r = rows[i]
        M = met.mcds[(r['dat'], r['mcd'])]
        pa = M['msgs'][r['msg']]['paras'][r['para']]
        if not isinstance(ko, str) or not ko.strip():
            errs.append(f'#{i}: 빈 번역'); continue
        if sorted(re.findall(r'\{btn:\d+\}', ko)) != sorted(re.findall(r'\{btn:\d+\}', r['ja'])):
            errs.append(f'#{i}: 버튼 태그 불일치 {re.findall(r"{btn:\d+}", r["ja"])} vs {re.findall(r"{btn:\d+}", ko)}')
        body = re.sub(r'\{btn:\d+\}', '', ko)
        if re.search(r'[{}]', body):
            errs.append(f'#{i}: 허용되지 않는 중괄호/태그: {ko!r}')
        if JP_RE.search(body) and not re.fullmatch(r'[^぀-ヿ㐀-鿿]*', body):
            bad = set(JP_RE.findall(body)) - set('・')
            if bad:
                errs.append(f'#{i}: 일본어/전각 문자 남음 {sorted(bad)}')
        odd = set(c for c in body if not ALLOWED_RE.match(c))
        if odd:
            errs.append(f'#{i}: 허용되지 않는 문자 {sorted(odd)}')
        lines = ko.split('\n')
        if len(lines) > len(pa['lines']):
            errs.append(f'#{i}: 줄 수 초과 {len(lines)} > {len(pa["lines"])}')
        widths = met.jp_line_widths(M, pa); mx = max(widths)
        limit = max(mx * 1.08, mx + met.hangul_w(pa['font']))
        for li, t in enumerate(lines):
            a = pa['lines'][min(li, len(pa['lines']) - 1)]['a']
            w = met.ko_line_width(t, pa['font'], a)
            if w > limit:
                errs.append(f'#{i}: {li + 1}번째 줄 폭 초과 {w:.0f}px > 한도 {limit:.0f}px (약 {int((w - limit) // met.hangul_w(pa["font"])) + 1}자 줄이기): {t!r}')
    if not quiet:
        print(f'[{name}] 번역 {len(outs)}/{len(src)}, 오류 {len(errs)}')
        for e in errs:
            print('  ' + e)
    return errs, outs


def merge():
    met = Metrics()
    allout = {}
    for p in sorted(glob.glob(os.path.join(WORK, 'in_*.json'))):
        name = os.path.basename(p)[3:-5]
        errs, outs = check(name, met, quiet=True)
        print(name, len(outs), 'errors', len(errs))
        allout.update(outs)
    wb = openpyxl.load_workbook(XLSX); ws = wb['text']
    n = 0
    for row in ws.iter_rows(min_row=2):
        if row[0].value in allout:
            row[8].value = allout[row[0].value]; n += 1
    wb.save(XLSX)
    print('merged', n)


if __name__ == '__main__':
    cmd = sys.argv[1]
    if cmd == 'split': split()
    elif cmd == 'check':
        met = Metrics()
        names = [os.path.basename(p)[3:-5] for p in sorted(glob.glob(os.path.join(WORK, 'in_*.json')))] if sys.argv[2] == 'all' else [sys.argv[2]]
        for n in names: check(n, met)
    elif cmd == 'merge': merge()
