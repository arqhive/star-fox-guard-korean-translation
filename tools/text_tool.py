"""스타폭스 가드 MCD 텍스트 추출/검증.
  python text_tool.py extract   -> translation/sfg_text.xlsx
  python text_tool.py roundtrip -> xlsx 원문을 다시 words로 인코딩해 원본 MCD와 바이트 비교
"""
import sys, glob, os, json
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from dat_lib import read_dat
from mcd_lib import parse_mcd, build_mcd, words_to_text, text_to_words

EXTRACT = os.path.join(os.path.dirname(__file__), '..', 'extract')
XLSX = os.path.join(os.path.dirname(__file__), '..', 'translation', 'sfg_text.xlsx')
HEAD = ['ID', 'DAT', 'MCD', 'Msg', 'Para', 'Font', 'Event', '일본어 원문', '한국어 번역', '메모']


def dat_files():
    return sorted(glob.glob(os.path.join(EXTRACT, 'ui_*.dat'))) + [os.path.join(EXTRACT, 'core_coreui.dat')]


def iter_mcds():
    for p in dat_files():
        dat = os.path.basename(p).split('_', 1)[1]
        for e in read_dat(open(p, 'rb').read()):
            if e['name'].endswith('.mcd'):
                yield p, dat, e['name'], e['data']


def para_text(pa, syms):
    return '\n'.join(words_to_text(l['words'], syms, pa['font']) for l in pa['lines'])


def extract():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'text'
    ws.append(HEAD)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBF7')
    n = 0
    for p, dat, mcd, data in iter_mcds():
        M = parse_mcd(data)
        ev = {}
        for e in M['events']: ev.setdefault(e['idx'], e['name'].split(b'\0')[0].decode())
        for mi, msg in enumerate(M['msgs']):
            for pi, pa in enumerate(msg['paras']):
                n += 1
                ws.append([n, dat, mcd, mi, pi, pa['font'], ev.get(mi, ''), para_text(pa, M['syms']), '', ''])
    for col, w in zip('ABCDEFGHIJ', [7, 22, 24, 6, 6, 6, 30, 50, 50, 20]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[7:9]: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'H2'
    ws.auto_filter.ref = ws.dimensions
    g = wb.create_sheet('안내')
    for line in [
        ['표기', '의미'],
        ['줄바꿈', '게임 내 줄바꿈 (문단 안의 줄)'],
        ['{btn:n}', '버튼 아이콘 (n = 아이콘 번호) - 번역문에 그대로 유지'],
        ['{k:n}', '바로 앞 글자의 자간 보정(픽셀). 원문 폰트용 미세조정이라 번역문에서는 생략 가능'],
        ['{{ }}', '실제 중괄호 문자'],
        ['한국어 번역 칸이 비어 있으면', '원문 유지'],
    ]:
        g.append(line)
    g.column_dimensions['A'].width = 28; g.column_dimensions['B'].width = 70
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    print('rows', n, '->', XLSX)


def roundtrip():
    ws = openpyxl.load_workbook(XLSX).active
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        rows[(r[1], r[2], r[3], r[4])] = r[7] or ''
    ok = bad = 0
    for p, dat, mcd, data in iter_mcds():
        M = parse_mcd(data)
        symidx = {(s[0], s[1]): i for i, s in enumerate(M['syms'])}
        for mi, msg in enumerate(M['msgs']):
            for pi, pa in enumerate(msg['paras']):
                lines = rows.pop((dat, mcd, mi, pi)).split('\n')
                assert len(lines) == len(pa['lines']), (mcd, mi, pi)
                for l, t in zip(pa['lines'], lines):
                    l['words'] = text_to_words(t, symidx, pa['font'])
        if build_mcd(M) == data: ok += 1
        else: bad += 1; print('MISMATCH', dat, mcd)
    print('mcd identical', ok, 'mismatch', bad, 'unused xlsx rows', len(rows))


if __name__ == '__main__':
    {'extract': extract, 'roundtrip': roundtrip}[sys.argv[1]]()
